"""Carteiras (ADR 0001) — Frente 3, leva 4 (09/10/2026).

Saiu de `app/main.py` sem mudar comportamento: quem recebe carteira, as
Atribuições do Operador e as Exceções do Gestor, a importação por planilha.
A regra de quem pode gerir quem (`require_admin_ou_gestor`, `_exigir_alcance`)
mora em `app/alcance.py`, porque Usuários e Instalador também a usam.
"""

from __future__ import annotations

import csv
import io
import logging
from datetime import datetime, timedelta, timezone
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from pydantic import BaseModel

from app import auth, cert_installer, nomes, permissoes
from app.alcance import _documentos_ao_alcance, _exigir_alcance, _papel_do_alvo, require_admin_ou_gestor
from app.comum import (
    ERRO_INTERNO_VEJA_LOG,
    ERRO_LINHAS_DEMAIS,
    MAX_LINHAS_IMPORT,
    _documento_formatado,
    _ler_upload_limitado,
    _norm_header,
)
from app.sessao import _user_id_da_sessao, require_modulo

logger = logging.getLogger("app.main")
router = APIRouter()


def _exigir_documentos_no_inventario(documentos: List[str]) -> None:
    """Atribuição e Exceção são sobre Documentos do Inventário (GLOSSARY).

    A planilha já recusava documento fora do inventário; a tela e a API
    aceitavam qualquer número. Um CNPJ digitado errado viraria uma linha de
    carteira que nunca casa com certificado nenhum — e ninguém saberia.
    Retirar continua livre: documento que saiu do inventário precisa poder
    sair da carteira também.
    """
    universo = {d["documento"] for d in cert_installer.universo_de_documentos()}
    fora = sorted({
        cert_installer.so_digitos(d) for d in documentos
        if cert_installer.so_digitos(d) and cert_installer.so_digitos(d) not in universo
    })
    if fora:
        raise HTTPException(status_code=422, detail="Fora do inventário: " + ", ".join(fora))


def _exigir_documentos_ao_alcance(token: auth.TokenData, documentos: List[str]) -> None:
    """Gestor só atribui o que está no próprio Alcance (ADR 0001).

    Uma Exceção registrada para o Gestor também o impede de dar aquele
    Documento aos seus Operadores — senão ele atribuiria X a um Operador e
    pediria para instalar na estação dele. Só o Administrador o faz.
    """
    alcance = _documentos_ao_alcance(token)
    if alcance is None:
        return
    negados = sorted({
        cert_installer.so_digitos(d) for d in documentos
        if cert_installer.so_digitos(d) and cert_installer.so_digitos(d) not in alcance
    })
    if negados:
        raise HTTPException(
            status_code=403,
            detail="Fora do seu alcance, só o administrador atribui: " + ", ".join(negados),
        )


class CarteiraRequest(BaseModel):
    user_id: str
    documentos: List[str]


# Modulo `carteiras` na matriz desde 20/08, com os DOIS eixos declarados: a
# matriz diz se o papel alcanca (e se so ve ou tambem monta), e
# `require_admin_ou_gestor` diz de QUEM. Um lider do Fiscal com "Ver e editar"
# continua sem tocar na carteira de alguem do Contabil.
@router.get("/api/carteira/operadores", dependencies=[Depends(require_modulo("carteiras"))])
def listar_operadores(
    token: auth.TokenData = Depends(require_admin_ou_gestor),
) -> dict:
    """
    Quem pode receber carteira, com quantos documentos cada um já tem.

    Rota própria em vez de `/api/users`: aquela é de admin e devolve a linha
    inteira do usuário. O gestor precisa montar carteira **sem** poder
    administrar contas, e não tem por que ver o hash de senha de ninguém.
    """
    from app.settings_state import _banco

    sb = _banco()
    if not sb:
        raise HTTPException(status_code=503, detail="Banco não configurado")
    try:
        us = sb.table("users").select(
            "id, email, full_name, role, ativo, departamento_id"
        ).execute().data or []
        cart = sb.table("carteira").select("user_id").execute().data or []
        exc = sb.table(cert_installer.TABELA_EXCECOES).select("user_id").execute().data or []
    except Exception:
        logger.exception("Falha ao listar operadores")
        raise HTTPException(status_code=503, detail="Não foi possível listar os operadores.")

    from collections import Counter

    # O líder vê apenas quem ele alcança. Mostrar a lista inteira e recusar
    # depois seria a pior combinação: ele monta a carteira de alguém do outro
    # setor, clica em salvar, e só aí descobre — sem entender o critério.
    #
    # Filtrado aqui e conferido de novo em cada rota que age: esta é a tela, e
    # tela não é barreira. Quem chamar a API direto com outro `user_id` esbarra
    # em `_exigir_alcance`.
    eu = _user_id_da_sessao(token) or ""
    papel = (token.role or "").strip().lower()
    if papel in cert_installer.PAPEIS_COM_ALCANCE_TOTAL:
        visiveis = us
    else:
        try:
            meus = cert_installer.departamentos_que_lidera(eu)
        except cert_installer.AlcanceIndisponivel:
            raise HTTPException(
                status_code=503,
                detail="Não foi possível verificar seus departamentos. Tente de novo.",
            )
        # Gestor edita só Atribuições: Operadores dos departamentos que
        # lidera. Nem ele mesmo (a carteira dele é de Exceções, do
        # administrador), nem outro Gestor.
        visiveis = [
            u for u in us
            if (u.get("role") or "").strip().lower() == cert_installer.PAPEL_OPERADOR
            and u.get("departamento_id") and str(u["departamento_id"]) in meus
        ]

    quantos = Counter(str(c.get("user_id")) for c in cart)
    excecoes = Counter(str(c.get("user_id")) for c in exc)
    # O tamanho do inventário só interessa se houver Gestor na lista: a
    # carteira dele é "tudo menos N", e "tudo" é isto.
    universo = len(cert_installer.universo_de_documentos()) if any(
        (u.get("role") or "").strip().lower() == cert_installer.PAPEL_GESTOR for u in visiveis
    ) else 0
    from app import texto as _texto
    operadores = []
    for u in visiveis:
        uid = str(u.get("id"))
        papel_u = (u.get("role") or "").strip().lower()
        n_atrib = quantos.get(uid, 0)
        n_exc = excecoes.get(uid, 0)
        # Operador depende de Atribuições; Gestor tem tudo menos Exceções; o
        # administrador tem alcance sem carteira (ADR 0001).
        if papel_u == cert_installer.PAPEL_GESTOR:
            tipo, n = "excecoes", max(universo - n_exc, 0)
            clientes = ("todos os clientes" if not n_exc else f"todos menos {_texto.plural(n_exc, 'exceção', 'exceções')}")
        elif papel_u == cert_installer.PAPEL_OPERADOR:
            tipo, n = "atribuicoes", n_atrib
            clientes = _texto.plural(n, "cliente")
        else:
            tipo, n = "total", universo
            clientes = "alcance total"
        depende_de_carteira = tipo == "atribuicoes"
        operadores.append(
            {
                "id": uid,
                "email": u.get("email"),
                "full_name": u.get("full_name"),
                # Só na exibição: "irla" → "Irla", caixa alta → título. O dado
                # gravado não muda (a correção do cadastro é em Usuários).
                "nome_exibicao": nomes.nome_pessoa(u.get("full_name")) or (u.get("email") or ""),
                "role": u.get("role"),
                "ativo": auth.conta_ativa(u),
                "departamento_id": u.get("departamento_id"),
                "documentos": n,
                "tipo_carteira": tipo,
                "excecoes": n_exc,
                "depende_de_carteira": depende_de_carteira,
                "sem_carteira": n == 0 and depende_de_carteira,
                "textos": {"clientes": clientes},
            }
        )
    # Sem carteira primeiro (é o único estado que pede ação; o Dashboard
    # aponta para cá por isso), depois em ordem alfabética.
    operadores.sort(key=lambda o: (0 if o["sem_carteira"] else 1, (o["nome_exibicao"] or "").lower()))

    # Resumo da lista, com a mesma regra da tela: Operadores ativos, mais
    # inativos que ainda tenham carteira a limpar. Gestores aparecem para o
    # administrador mas não contam como "operador" no resumo.
    na_lista = [o for o in operadores if o["depende_de_carteira"] and (o["ativo"] or o["documentos"] > 0)]
    sem = sum(1 for o in na_lista if o["sem_carteira"])
    total = len(na_lista)
    if sem:
        resumo_txt = _texto.plural(total, "operador", "operadores") + " · " + str(sem) + " sem carteira"
    else:
        resumo_txt = _texto.plural(total, "operador", "operadores") + (", todos com carteira" if total else "")
    return {
        "operadores": operadores,
        "resumo": {"operadores": total, "sem_carteira": sem, "texto": resumo_txt},
    }


@router.get("/api/carteira/documentos", dependencies=[Depends(require_modulo("carteiras"))])
def listar_documentos_atribuiveis(
    q: Optional[str] = Query(None, max_length=120),
    limite: int = Query(500, ge=1, le=2000),
    token: auth.TokenData = Depends(require_admin_ou_gestor),
) -> dict:
    """
    Universo de documentos para atribuir, filtrável por nome ou número.

    Recortado pelo alcance de quem pergunta (01/10/2026): o Gestor só atribui
    o que alcança, então oferecer-lhe o inventário inteiro era convidá-lo a
    selecionar o que o servidor recusaria — e um "Todos" falhava inteiro.
    Para o administrador, `None` = tudo.

    O teto era 500 e havia 491 clientes — a um cadastro de distância de
    truncar em silêncio, que é como a curva de vencimento perdeu 29
    certificados em 15/08. Subiu para 2000; a tela pede a lista inteira de
    uma vez (medido: 491 documentos = 33 KB em ~375 ms) e monta os dois
    painéis no cliente, então filtrar no servidor virou opcional.

    `total` continua vindo separado de `documentos` justamente para a tela
    poder dizer quando a lista foi cortada, em vez de parecer completa.
    """
    todos = cert_installer.universo_de_documentos()
    alcance = _documentos_ao_alcance(token)
    if alcance is not None:
        todos = [d for d in todos if d["documento"] in alcance]
    termo = (q or "").strip().lower()
    if termo:
        digitos = cert_installer.so_digitos(termo)
        todos = [
            d for d in todos
            if termo in (d["nome"] or "").lower()
            or (digitos and digitos in d["documento"])
        ]
    # Máscara no servidor, como nas outras telas: a regra mora num lugar só.
    saida = [{**d, "documento_formatado": _documento_formatado(d["documento"])} for d in todos[:limite]]
    return {"total": len(todos), "documentos": saida}


@router.get("/api/carteira/{user_id}", dependencies=[Depends(require_modulo("carteiras"))])
def obter_carteira(
    user_id: str,
    token: auth.TokenData = Depends(require_admin_ou_gestor),
) -> dict:
    """
    Documentos que este operador pode instalar, com a trilha de atribuição.

    A dependência saiu do decorador e virou parâmetro porque agora o token é
    USADO: `_exigir_alcance` precisa saber quem está perguntando. A trilha diz
    quem liberou o quê — informação de dentro do setor.
    """
    _exigir_alcance(token, user_id)
    nomes = {d["documento"]: d["nome"] for d in cert_installer.universo_de_documentos()}

    if _papel_do_alvo(user_id) == cert_installer.PAPEL_GESTOR:
        # Carteira de Exceções (ADR 0001): "na carteira" é o inventário menos
        # o que o administrador retirou; "disponíveis" são as Exceções. Os dois
        # painéis da tela continuam os mesmos — só o sinal inverte.
        try:
            excecoes = cert_installer.detalhar_excecoes(user_id)
        except cert_installer.CarteiraIndisponivel as e:
            raise HTTPException(status_code=503, detail=str(e))
        fora = {l["documento"]: l for l in excecoes}
        return {
            "user_id": user_id,
            "tipo": "excecoes",
            "documentos": sorted(d for d in nomes if d not in fora),
            "itens": [
                {"documento": d, "nome": nomes[d], "no_inventario": True,
                 "atribuido_por": None, "atribuido_em": None, "origem": "papel"}
                for d in sorted(nomes) if d not in fora
            ],
            "excecoes": [
                {"documento": l["documento"], "nome": nomes.get(l["documento"], ""),
                 "registrado_por": l.get("registrado_por_email"), "registrado_em": l.get("registrado_em")}
                for l in excecoes
            ],
        }

    try:
        linhas = cert_installer.detalhar_carteira(user_id)
    except cert_installer.CarteiraIndisponivel as e:
        raise HTTPException(status_code=503, detail=str(e))

    return {
        "user_id": user_id,
        "tipo": "atribuicoes",
        "documentos": sorted(l["documento"] for l in linhas),
        "itens": [
            {
                "documento": l["documento"],
                # Documento atribuído que não está mais no inventário continua
                # na carteira: a atribuição é uma decisão, e sumir com ela
                # esconderia que a pessoa tem acesso a algo que voltou depois.
                "nome": nomes.get(l["documento"], ""),
                "no_inventario": l["documento"] in nomes,
                "atribuido_por": l.get("atribuido_por_email"),
                "atribuido_em": l.get("atribuido_em"),
            }
            for l in linhas
        ],
    }


@router.get("/api/carteira/{user_id}/instalacoes", dependencies=[Depends(require_modulo("carteiras"))])
def instalacoes_da_pessoa(
    user_id: str,
    dias: int = Query(365, ge=1, le=365),
    token: auth.TokenData = Depends(require_admin_ou_gestor),
) -> dict:
    """
    As tentativas de instalação de quem está com a carteira aberta.

    Até 30/09 a tela chamava `/api/cert-installer/trilha`, do módulo
    Instalador, que o Gestor não alcança — ele via "Disponível para
    administradores". Decisão C3 (01/10/2026): a pergunta "o fulano instalou
    o que eu atribuí?" é do Gestor, sobre os Operadores dos departamentos que
    lidera. O recorte é o mesmo da carteira (`_exigir_alcance`), e o IP de
    origem só vai para o administrador.
    """
    _exigir_alcance(token, user_id)
    sb = cert_installer._banco()
    if not sb:
        raise HTTPException(status_code=503, detail="Banco não configurado")
    try:
        linhas = sb.table("users").select("email").eq("id", user_id).limit(1).execute().data or []
    except Exception:  # noqa: BLE001
        logger.exception("Falha ao ler o e-mail de %s", user_id)
        raise HTTPException(status_code=503, detail="Não foi possível ler a conta. Tente de novo.")
    if not linhas:
        raise HTTPException(status_code=404, detail="Pessoa não encontrada.")
    email_alvo = str(linhas[0].get("email") or "").strip().lower()
    desde = (datetime.now(timezone.utc) - timedelta(days=dias)).isoformat()
    cadeias = cert_installer.cadeias_de_instalacao(limite=200, desde=desde, user_email=email_alvo)
    if (token.role or "").strip().lower() not in cert_installer.PAPEIS_COM_ALCANCE_TOTAL:
        for c in cadeias:
            c.pop("client_ip", None)
    return {"dias": dias, "resumo": cert_installer.resumo_das_cadeias(cadeias), "cadeias": cadeias}


@router.post("/api/carteira", dependencies=[Depends(require_modulo("carteiras", permissoes.NIVEL_EDITAR))])
def atribuir_carteira(
    body: CarteiraRequest,
    token: auth.TokenData = Depends(require_admin_ou_gestor),
) -> dict:
    """
    Acrescenta documentos à carteira de um operador.

    Quem atribui fica registrado — e-mail inclusive, não só o UUID: essa
    trilha é a única forma de reconstruir o que houve se uma conta de gestor
    for comprometida; guardar só o UUID a perderia no dia em que a conta fosse
    apagada.

    Alvo Gestor (ADR 0001): "liberar" um documento é tirar a Exceção. Só o
    administrador chega aqui com alvo Gestor — `pode_gerir` barra os demais.
    """
    _exigir_alcance(token, body.user_id)
    _exigir_documentos_no_inventario(body.documentos)
    if _papel_do_alvo(body.user_id) == cert_installer.PAPEL_GESTOR:
        try:
            for doc in body.documentos:
                cert_installer.remover_excecao(body.user_id, doc)
        except Exception:
            logger.exception("Erro ao remover exceção")
            raise HTTPException(status_code=500, detail="Erro interno ao devolver o cliente ao gestor")
        return {"status": "ok", "gravados": len(body.documentos), "tipo": "excecoes"}
    _exigir_documentos_ao_alcance(token, body.documentos)
    try:
        gravados = cert_installer.atribuir_carteira(
            user_id=body.user_id,
            documentos=body.documentos,
            atribuido_por=_user_id_da_sessao(token),
            atribuido_por_email=token.email or "desconhecido",
        )
        return {"status": "ok", "gravados": gravados, "tipo": "atribuicoes"}
    except RuntimeError as e:
        # O texto vinha do cofre/banco e podia trazer nome de chave de
        # ambiente ou de tabela (achado #35). Fica no log, com a rota.
        logger.exception("Operação recusada pelo cofre ou pelo banco: %s", e)
        raise HTTPException(status_code=500, detail=ERRO_INTERNO_VEJA_LOG)
    except Exception:
        logger.exception("Erro ao atribuir carteira")
        raise HTTPException(status_code=500, detail="Erro interno ao atribuir carteira")


def _linhas_da_planilha(nome: str, raw: bytes) -> List[Dict[str, str]]:
    """
    Lê .csv ou .xlsx e devolve as linhas como dicionários de cabeçalho→valor.

    O .xlsx entra porque é o que sai do Excel sem passo extra — pedir "salve
    como CSV" antes de cada importação é exatamente o trabalho manual que
    esta rota existe para tirar. O `openpyxl` só é importado aqui: quem nunca
    importa planilha não paga o custo no cold start da Vercel.
    """
    if nome.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover - depende do ambiente de deploy
            raise HTTPException(
                status_code=503,
                detail="Leitura de .xlsx indisponível no servidor. Envie o arquivo como .csv.",
            )
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            # `max_row` limita a LEITURA: um .xlsx de 5 MB com dimensão gigante
            # expandiria para GB em RAM se materializado inteiro (achado #11).
            # Lê uma linha além do teto só para saber que ele foi passado.
            linhas = list(ws.iter_rows(values_only=True, max_row=MAX_LINHAS_IMPORT + 2))
        except Exception:
            raise HTTPException(status_code=422, detail="Não consegui ler a planilha .xlsx.")
        finally:
            try:
                wb.close()
            except Exception:
                pass
        if not linhas:
            raise HTTPException(status_code=422, detail="Planilha vazia.")
        if len(linhas) - 1 > MAX_LINHAS_IMPORT:
            raise HTTPException(status_code=413, detail=ERRO_LINHAS_DEMAIS)
        cabecalho = [str(c or "").strip() for c in linhas[0]]
        return [
            {cabecalho[i]: ("" if v is None else str(v).strip())
             for i, v in enumerate(linha) if i < len(cabecalho)}
            for linha in linhas[1:]
        ]

    # CSV: mesmo tratamento do import de usuários — BOM do Excel e separador
    # `;` do pt-BR são a regra, não a exceção.
    texto = raw.decode("utf-8-sig", errors="replace")
    try:
        delim = csv.Sniffer().sniff(texto[:2048], delimiters=",;").delimiter
    except csv.Error:
        delim = ";"
    leitor = csv.DictReader(io.StringIO(texto), delimiter=delim)
    if not leitor.fieldnames:
        raise HTTPException(status_code=422, detail="CSV sem cabeçalho.")
    import itertools

    linhas = [{(k or ""): (v or "") for k, v in linha.items()}
              for linha in itertools.islice(leitor, MAX_LINHAS_IMPORT + 1)]
    if len(linhas) > MAX_LINHAS_IMPORT:
        raise HTTPException(status_code=413, detail=ERRO_LINHAS_DEMAIS)
    return linhas


@router.post("/api/carteira/importar", dependencies=[Depends(require_modulo("carteiras", permissoes.NIVEL_EDITAR))])
async def importar_carteiras(
    request: Request,
    file: UploadFile = File(...),
    token: auth.TokenData = Depends(require_admin_ou_gestor),
) -> dict:
    """
    Atribui carteiras em massa a partir de uma planilha de e-mail + documento.

    Montar carteira clicando cliente a cliente não escala: quem recebe uma
    lista pronta da operação copiava CNPJ por CNPJ na mão.

    O ALCANCE É CONFERIDO POR PESSOA, e não uma vez para o arquivo: o gestor
    importa só para quem ele lidera, exatamente como na tela. Uma linha fora
    do alcance vira erro DAQUELA linha — recusar o arquivo inteiro por causa
    de uma pessoa faria o gestor perder o trabalho já correto.

    Nada é gravado até o arquivo inteiro ser lido e validado: um arquivo com
    metade das linhas erradas deixaria a carteira em estado parcial, e o
    operador não teria como saber o que entrou.
    """
    nome = (file.filename or "").lower()
    if not (nome.endswith(".csv") or nome.endswith(".xlsx")):
        raise HTTPException(
            status_code=422,
            detail="Formato inválido. Envie a planilha em .xlsx ou .csv.",
        )

    raw = await _ler_upload_limitado(request, file)
    if not raw:
        raise HTTPException(status_code=422, detail="Arquivo vazio.")
    # Executável disfarçado. O `PK` do zip é legítimo aqui — todo .xlsx começa
    # com ele —, então a checagem é por extensão declarada.
    if raw.startswith(b"MZ") or raw.startswith(b"\x7fELF") or raw.startswith(b"%PDF"):
        raise HTTPException(status_code=422, detail="Conteúdo do arquivo suspeito.")
    if nome.endswith(".xlsx") and not raw.startswith(b"PK"):
        raise HTTPException(status_code=422, detail="Isto não é um .xlsx válido.")

    linhas = _linhas_da_planilha(nome, raw)
    if not linhas:
        raise HTTPException(status_code=422, detail="A planilha não tem nenhuma linha de dados.")

    mapa = {_norm_header(h): h for linha in linhas[:1] for h in linha}

    def coluna(*aliases: str) -> Optional[str]:
        for a in aliases:
            k = mapa.get(_norm_header(a))
            if k:
                return k
        return None

    col_email = coluna("email", "e-mail", "colaborador", "email do colaborador", "usuario")
    col_doc = coluna("documento", "cnpj", "cpf", "cnpj/cpf", "cnpj da empresa", "cliente")
    if not col_email or not col_doc:
        raise HTTPException(
            status_code=422,
            detail="A planilha precisa de duas colunas: e-mail do colaborador e CNPJ/CPF do cliente.",
        )

    from app.settings_state import _banco

    sb = _banco()
    if not sb:
        raise HTTPException(status_code=503, detail="Banco não configurado.")
    try:
        contas = sb.table("users").select("id, email, role, ativo").execute().data or []
    except Exception:
        logger.exception("Falha ao ler usuários na importação de carteiras")
        raise HTTPException(status_code=503, detail="Não foi possível ler os usuários.")
    por_email = {str(u.get("email") or "").strip().lower(): str(u.get("id")) for u in contas}
    papel_de = {str(u.get("id")): (u.get("role") or "").strip().lower() for u in contas}
    inativa = {str(u.get("id")) for u in contas if not auth.conta_ativa(u)}

    universo = {d["documento"] for d in cert_installer.universo_de_documentos()}
    # Gestor só atribui o que alcança (ADR 0001); `None` é o administrador.
    meu_alcance = _documentos_ao_alcance(token)

    erros: List[Dict[str, Any]] = []
    por_usuario: Dict[str, set] = {}
    alcance: Dict[str, bool] = {}
    ator = _user_id_da_sessao(token) or ""

    for i, linha in enumerate(linhas, start=2):  # 1 é o cabeçalho
        email = (linha.get(col_email) or "").strip().lower()
        doc = cert_installer.so_digitos(linha.get(col_doc) or "")
        if not email and not doc:
            continue  # linha em branco no fim da planilha
        if not email or not doc:
            erros.append({"linha": i, "motivo": "Falta o e-mail ou o CNPJ/CPF."})
            continue

        user_id = por_email.get(email)
        if not user_id:
            # A linha identifica o registro na planilha de quem importou; o
            # endereço não precisa voltar na resposta (achado #52).
            erros.append({"linha": i, "motivo": "Não existe usuário com esse e-mail."})
            continue

        if user_id not in alcance:
            try:
                alcance[user_id] = cert_installer.pode_gerir(ator, token.role or "", user_id)
            except cert_installer.AlcanceIndisponivel:
                raise HTTPException(
                    status_code=503,
                    detail="Não foi possível verificar seu alcance. Tente de novo.",
                )
        if papel_de.get(user_id) != cert_installer.PAPEL_OPERADOR:
            # Planilha é de Atribuições. A carteira de um Gestor é de Exceções
            # e se edita na tela, pelo administrador.
            erros.append({"linha": i, "motivo": f"{email} não é Operador; a planilha só atribui a Operadores."})
            continue
        if user_id in inativa:
            # Desativar esvazia a carteira; atribuir a quem está desativado
            # gravaria o que a reativação não devolve.
            erros.append({"linha": i, "motivo": f"{email} está desativado; reative a conta antes de atribuir."})
            continue
        if not alcance[user_id]:
            erros.append({"linha": i, "motivo": f"{email} não está em um departamento que você lidera."})
            continue

        if doc not in universo:
            erros.append({"linha": i, "motivo": f"O documento {doc} não está no inventário."})
            continue
        if meu_alcance is not None and doc not in meu_alcance:
            erros.append({"linha": i, "motivo": f"O documento {doc} está fora do seu alcance; só o administrador o libera."})
            continue

        por_usuario.setdefault(user_id, set()).add(doc)

    atribuidos = 0
    pessoas = 0
    for user_id, docs in por_usuario.items():
        try:
            atribuidos += cert_installer.atribuir_carteira(
                user_id=user_id,
                documentos=sorted(docs),
                atribuido_por=ator,
                atribuido_por_email=token.email or "desconhecido",
            )
            pessoas += 1
        except Exception:
            logger.exception("Falha ao gravar carteira importada")
            erros.append({"linha": 0, "motivo": "Falha ao gravar a carteira de um dos operadores."})

    return {
        "status": "ok",
        "atribuidos": atribuidos,
        "pessoas": pessoas,
        "linhas_lidas": len(linhas),
        "erros": erros,
    }


@router.delete("/api/carteira/{user_id}/{documento}", dependencies=[Depends(require_modulo("carteiras", permissoes.NIVEL_EDITAR))])
def remover_carteira(
    user_id: str,
    documento: str,
    token: auth.TokenData = Depends(require_admin_ou_gestor),
) -> dict:
    _exigir_alcance(token, user_id)
    if _papel_do_alvo(user_id) == cert_installer.PAPEL_GESTOR:
        # Tirar um cliente da carteira de um Gestor é registrar uma Exceção
        # (ADR 0001). Só o administrador chega aqui com alvo Gestor.
        _exigir_documentos_no_inventario([documento])
        try:
            cert_installer.registrar_excecoes(
                user_id, [documento], _user_id_da_sessao(token), token.email or "desconhecido",
            )
            return {"status": "ok", "user_id": user_id, "documento": documento, "tipo": "excecoes"}
        except Exception:
            logger.exception("Erro ao registrar exceção")
            raise HTTPException(status_code=500, detail="Erro interno ao registrar a exceção")
    try:
        cert_installer.remover_da_carteira(user_id, documento)
        return {"status": "ok", "user_id": user_id, "documento": documento, "tipo": "atribuicoes"}
    except RuntimeError as e:
        # O texto vinha do cofre/banco e podia trazer nome de chave de
        # ambiente ou de tabela (achado #35). Fica no log, com a rota.
        logger.exception("Operação recusada pelo cofre ou pelo banco: %s", e)
        raise HTTPException(status_code=500, detail=ERRO_INTERNO_VEJA_LOG)
    except Exception:
        logger.exception("Erro ao remover da carteira")
        raise HTTPException(status_code=500, detail="Erro interno ao remover da carteira")
